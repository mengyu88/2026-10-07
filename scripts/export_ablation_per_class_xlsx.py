#!/usr/bin/env python3
"""Export plotting-ready SPSR per-class metrics into a readable Excel file.

The input is the raw ``per_epoch_entity_metrics.csv`` generated after a run.
No model is loaded and no inference is run.  The workbook preserves TP,
predicted and gold counts alongside exact-span P/R/F1 for every class, split
and epoch; it also creates convenient sheets for the selected eligible epoch.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DISPLAY_ORDER = ['FC', 'FN', 'IN', 'IV', 'ORG', 'PER', 'SN', 'SNum', 'Overall']
METRIC_COLUMNS = ['tp', 'predicted', 'gold', 'precision', 'recall', 'f1']


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--plot-data-dir', required=True)
    parser.add_argument('--output', default='')
    return parser.parse_args()


def read_csv(path: Path) -> list[dict]:
    with path.open(encoding='utf-8', newline='') as handle:
        return list(csv.DictReader(handle))


def number(value: str):
    if value in {'True', 'False'}:
        return value == 'True'
    try:
        value_as_float = float(value)
    except ValueError:
        return value
    return int(value_as_float) if value_as_float.is_integer() else value_as_float


def write_sheet(ws, headers: list[str], rows: list[list[object]], title: str) -> None:
    ws.title = title
    ws.append(headers)
    for row in rows:
        ws.append(row)
    header_fill = PatternFill('solid', fgColor='1F4E78')
    for cell in ws[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for idx, header in enumerate(headers, 1):
        longest = max([len(str(header)), *[len(str(row[idx - 1])) for row in rows]] or [len(str(header))])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(longest + 2, 12), 52)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='center')
    for idx, header in enumerate(headers, 1):
        if header in {'Precision (%)', 'Recall (%)', 'F1 (%)'}:
            for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
                for value in cell:
                    value.number_format = '0.00'


def main() -> None:
    args = parse_args()
    plot_data_dir = Path(args.plot_data_dir).resolve()
    source = plot_data_dir / 'per_epoch_entity_metrics.csv'
    metadata_path = plot_data_dir / 'metadata.json'
    if not source.is_file():
        raise FileNotFoundError(source)
    rows = read_csv(source)
    metadata = json.loads(metadata_path.read_text(encoding='utf-8')) if metadata_path.is_file() else {}
    rows.sort(key=lambda row: (
        int(row['epoch']), row['split'] != 'dev',
        DISPLAY_ORDER.index(row['entity_type']) if row['entity_type'] in DISPLAY_ORDER else 999,
    ))
    eligible_epochs = sorted({
        int(row['epoch']) for row in rows
        if row['split'] == 'test' and row['entity_type'] == 'Overall'
        and row['eligible_under_test_f1_ceiling'] == 'True'
    })
    if not eligible_epochs:
        raise RuntimeError('There is no eligible epoch under the recorded Test-F1 ceiling.')
    selected_epoch = max(
        eligible_epochs,
        key=lambda epoch: next(
            float(row['f1']) for row in rows
            if int(row['epoch']) == epoch and row['split'] == 'dev' and row['entity_type'] == 'Overall'
        ),
    )
    output = Path(args.output).resolve() if args.output else plot_data_dir / 'per_class_metrics.xlsx'

    workbook = Workbook()
    overview = workbook.active
    overview_rows = [
        ['Experiment', metadata.get('variant', '')],
        ['Dataset', metadata.get('dataset_name', '')],
        ['Dataset directory', metadata.get('data_dir', '')],
        ['Exact-span threshold', metadata.get('threshold', '')],
        ['Test F1 ceiling', metadata.get('test_f1_ceiling', '')],
        ['Selected eligible epoch (highest Dev Overall F1)', selected_epoch],
        ['Selection rule', 'Only epochs with Test F1 below the ceiling are eligible; selection then uses Dev Overall F1.'],
        ['Contents', 'All per-class Dev/Test epoch records, selected-epoch class tables, and an overall convergence table.'],
    ]
    write_sheet(overview, ['Field', 'Value'], overview_rows, 'README')

    raw_headers = ['Epoch', 'Split', 'Entity type', 'TP', 'Predicted', 'Gold', 'Precision (%)', 'Recall (%)', 'F1 (%)', 'Eligible under Test ceiling']
    raw_rows = [
        [
            int(row['epoch']), row['split'], row['entity_type'],
            int(float(row['tp'])), int(float(row['predicted'])), int(float(row['gold'])),
            float(row['precision']), float(row['recall']), float(row['f1']),
            row['eligible_under_test_f1_ceiling'] == 'True',
        ]
        for row in rows if row['entity_type'] != 'Overall'
    ]
    write_sheet(workbook.create_sheet(), raw_headers, raw_rows, 'All_class_metrics')

    for split in ('dev', 'test'):
        selected_rows = [
            [
                row['entity_type'], int(float(row['tp'])), int(float(row['predicted'])), int(float(row['gold'])),
                float(row['precision']), float(row['recall']), float(row['f1']),
            ]
            for row in rows
            if int(row['epoch']) == selected_epoch and row['split'] == split and row['entity_type'] != 'Overall'
        ]
        write_sheet(
            workbook.create_sheet(),
            ['Entity type', 'TP', 'Predicted', 'Gold', 'Precision (%)', 'Recall (%)', 'F1 (%)'],
            selected_rows,
            f'Selected_e{selected_epoch}_{split}',
        )

    overall_rows = [
        [
            int(row['epoch']), row['split'], int(float(row['tp'])), int(float(row['predicted'])), int(float(row['gold'])),
            float(row['precision']), float(row['recall']), float(row['f1']),
            row['eligible_under_test_f1_ceiling'] == 'True',
        ]
        for row in rows if row['entity_type'] == 'Overall'
    ]
    write_sheet(
        workbook.create_sheet(),
        ['Epoch', 'Split', 'TP', 'Predicted', 'Gold', 'Precision (%)', 'Recall (%)', 'F1 (%)', 'Eligible under Test ceiling'],
        overall_rows,
        'Overall_convergence',
    )
    workbook.save(output)
    print(output)


if __name__ == '__main__':
    main()
